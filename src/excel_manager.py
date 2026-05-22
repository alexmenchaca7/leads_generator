import logging
import os
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config import MASTER_FILE, RUNS_DIR, DROPDOWN_OPTIONS
from src.lead_scoring import calculate_score
from src.utils import generate_business_id, normalize_url

logger = logging.getLogger(__name__)

# Manual-only sheet logging contact history (system never writes data rows here)
CONTACTED_COLUMNS = [
    "business_id", "name", "phone", "contacted_date",
    "response", "next_action", "deal_status", "notes",
]

# Rows to cover with dropdowns so manually-added / future-scraped rows inherit them
_DROPDOWN_LAST_ROW = 5000

# ── Column order for raw_leads ─────────────────────────────────────────────────
COLUMNS = [
    "business_id",
    "name",
    "category",
    "address",
    "phone",
    "website",
    "rating",
    "reviews_count",
    "lat",
    "lng",
    "maps_url",
    "no_website",
    "lead_score",
    "priority",
    "website_status",
    "first_seen",
    "last_seen",
    "outreach_status",
    "contacted",
    "follow_up",
    "notes",
]

COLUMN_WIDTHS = {
    "business_id":    15,
    "name":           35,
    "category":       28,
    "address":        40,
    "phone":          18,
    "website":        38,
    "rating":          8,
    "reviews_count":  14,
    "lat":            14,
    "lng":            14,
    "maps_url":       55,
    "no_website":     12,
    "lead_score":     12,
    "priority":       10,
    "website_status": 16,
    "first_seen":     13,
    "last_seen":      13,
    "outreach_status":20,
    "contacted":      12,
    "follow_up":      13,
    "notes":          45,
}

# Styles
_HDR_FILL   = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HDR_FONT   = Font(color="FFFFFF", bold=True, size=10)
_HIGH_FILL  = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
_NO_WEB_FILL= PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

# Sheets that must NEVER be deleted or auto-populated
_PROTECTED_SHEETS = {"contacted"}

# Sheets that are auto-refreshed (can be rebuilt)
_AUTO_SHEETS = {"no_website_leads", "high_priority"}


def _safe_save(wb: Workbook, path: Path):
    """
    Save via a temp file + atomic rename so that:
    - A partial write never corrupts the master file.
    - If the target is locked by Excel, we fail fast with a clear message
      instead of silently corrupting the file.
    """
    tmp_fd, tmp_path = tempfile.mkstemp(dir=path.parent, suffix=".tmp.xlsx")
    os.close(tmp_fd)
    try:
        wb.save(tmp_path)
        # On Windows, os.replace handles the locked-file case:
        # if the target is open in Excel, this raises PermissionError here
        # (not after a partial write), so the original is always intact.
        os.replace(tmp_path, path)
    except PermissionError:
        os.unlink(tmp_path)
        raise PermissionError(
            f"\n\n  El archivo está abierto en Excel:\n  {path}\n\n"
            "  Cierra Excel y vuelve a correr el script.\n"
        )
    except Exception:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
        raise


def _write_headers(ws, columns: list[str] = COLUMNS):
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = col_name
        cell.font = _HDR_FONT
        cell.fill = _HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS.get(col_name, 15)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _create_new_workbook() -> Workbook:
    wb = Workbook()

    ws_raw = wb.active
    ws_raw.title = "raw_leads"
    _write_headers(ws_raw)
    _apply_dropdowns(ws_raw, COLUMNS)

    ws_no_web = wb.create_sheet("no_website_leads")
    _write_headers(ws_no_web)

    ws_high = wb.create_sheet("high_priority")
    _write_headers(ws_high)

    # contacted sheet — manual only, system never writes here
    ws_c = wb.create_sheet("contacted")
    _write_headers(ws_c, CONTACTED_COLUMNS)
    _apply_dropdowns(ws_c, CONTACTED_COLUMNS)

    return wb


def _apply_dropdowns(ws, columns: list[str]):
    """Attach list-validation dropdowns to any column in this sheet that has
    options defined in DROPDOWN_OPTIONS. Re-applied on every save, so we clear
    existing validations first to avoid accumulating duplicates across runs.
    """
    ws.data_validations.dataValidation = []  # reset

    for col_name, options in DROPDOWN_OPTIONS.items():
        if col_name not in columns:
            continue
        col_letter = get_column_letter(columns.index(col_name) + 1)
        # Excel list formula: comma-separated values wrapped in quotes.
        formula = '"' + ",".join(options) + '"'
        dv = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showErrorMessage=True,
        )
        dv.errorTitle = "Valor no permitido"
        dv.error      = "Elige uno de los valores de la lista."
        dv.promptTitle = col_name
        dv.prompt      = "Selecciona un valor de la lista."
        dv.add(f"{col_letter}2:{col_letter}{_DROPDOWN_LAST_ROW}")
        ws.add_data_validation(dv)


def _row_fill(no_website: bool, priority: str):
    """Background fill for a raw_leads row, or None for plain rows."""
    if no_website and priority == "high":
        return _HIGH_FILL
    if no_website:
        return _NO_WEB_FILL
    return None


def _business_to_row(business: dict, first_seen: str, last_seen: str) -> list:
    no_website = not bool((business.get("website") or "").strip())
    row = []
    for col in COLUMNS:
        if col == "no_website":
            row.append(no_website)
        elif col == "website_status":
            row.append("no_website" if no_website else "has_website")
        elif col == "first_seen":
            row.append(first_seen)
        elif col == "last_seen":
            row.append(last_seen)
        else:
            row.append(business.get(col) or "")
    return row


class ExcelManager:
    def __init__(self):
        MASTER_FILE.parent.mkdir(parents=True, exist_ok=True)
        RUNS_DIR.mkdir(parents=True, exist_ok=True)

    # ── Internal helpers ───────────────────────────────────────────────────────

    def _load_or_create(self) -> Workbook:
        if MASTER_FILE.exists():
            return load_workbook(str(MASTER_FILE))
        logger.info("Creating new workbook: %s", MASTER_FILE)
        return _create_new_workbook()

    def _ensure_sheets(self, wb: Workbook):
        """Make sure all required sheets exist (safe for existing workbooks)."""
        if "raw_leads" not in wb.sheetnames:
            ws = wb.create_sheet("raw_leads", 0)
            _write_headers(ws)
        for name in ("no_website_leads", "high_priority"):
            if name not in wb.sheetnames:
                ws = wb.create_sheet(name)
                _write_headers(ws)
        if "contacted" not in wb.sheetnames:
            ws = wb.create_sheet("contacted")
            _write_headers(ws, CONTACTED_COLUMNS)
            _apply_dropdowns(ws, CONTACTED_COLUMNS)

    # ── Deduplication ──────────────────────────────────────────────────────────

    def _load_existing(self) -> tuple[set[str], set[str]]:
        """Return (existing_ids, existing_normalized_urls)."""
        if not MASTER_FILE.exists():
            return set(), set()
        try:
            df = pd.read_excel(str(MASTER_FILE), sheet_name="raw_leads", dtype=str)
            ids  = set(df["business_id"].dropna().tolist())
            urls = set(
                df["maps_url"]
                .dropna()
                .apply(normalize_url)
                .tolist()
            )
            return ids, urls
        except Exception as exc:
            logger.warning("Could not read existing data for dedup: %s", exc)
            return set(), set()

    # ── Public API ─────────────────────────────────────────────────────────────

    def existing_keys(self) -> tuple[set[str], set[str]]:
        """Return (existing_ids, existing_normalized_urls) from the master file."""
        return self._load_existing()

    def _compact_raw(self, wb: Workbook):
        """Remove fully-empty rows from raw_leads and re-pack the data.

        When a user deletes rows by clearing their contents in Excel, openpyxl
        still counts those blank rows in ``max_row``. ``ws.append`` then writes
        new data *after* the gap, leaving empty cells in the middle. Re-packing
        the sheet on every save makes those gaps self-heal and keeps appends
        landing right below the last real row. All columns (including the manual
        outreach/notes ones the user edits) are preserved.
        """
        ws = wb["raw_leads"]
        ncols = len(COLUMNS)
        if ws.max_row < 2:
            return

        kept = [
            list(row)
            for row in ws.iter_rows(min_row=2, max_col=ncols, values_only=True)
            if any(v is not None and str(v).strip() != "" for v in row)
        ]
        if len(kept) == ws.max_row - 1:
            return  # already contiguous, nothing to do

        ws.delete_rows(2, ws.max_row - 1)

        web_idx  = COLUMNS.index("website")
        prio_idx = COLUMNS.index("priority")
        for row in kept:
            ws.append(row)
            no_website = not bool(str(row[web_idx] or "").strip())
            fill = _row_fill(no_website, row[prio_idx])
            if fill:
                rn = ws.max_row
                for col_idx in range(1, ncols + 1):
                    ws.cell(row=rn, column=col_idx).fill = fill

        ws.auto_filter.ref = ws.dimensions
        logger.info("Compacted raw_leads → %d rows (removed empty gaps)", len(kept))

    def _ensure_dropdowns(self, wb: Workbook):
        """(Re)apply the list-validation dropdowns to the editable sheets."""
        _apply_dropdowns(wb["raw_leads"], COLUMNS)
        if "contacted" in wb.sheetnames:
            _apply_dropdowns(wb["contacted"], CONTACTED_COLUMNS)

    def save_new_businesses(self, businesses: list[dict]) -> tuple[int, int]:
        """Append new businesses to raw_leads and refresh filtered sheets.

        Returns (new_count, duplicate_count).
        """
        wb = self._load_or_create()
        self._ensure_sheets(wb)
        self._compact_raw(wb)
        ws_raw = wb["raw_leads"]

        existing_ids, existing_urls = self._load_existing()
        now = datetime.now().strftime("%Y-%m-%d")
        new_count = dup_count = 0

        for biz in businesses:
            maps_url = normalize_url(biz.get("maps_url") or "")
            biz_id   = generate_business_id(
                maps_url=maps_url,
                name=biz.get("name", ""),
                address=biz.get("address", ""),
            )

            # Dedup check
            if (maps_url and maps_url in existing_urls) or biz_id in existing_ids:
                dup_count += 1
                continue

            # Enrich
            biz["business_id"] = biz_id
            biz["maps_url"]    = maps_url
            biz.update(calculate_score(biz))

            row     = _business_to_row(biz, first_seen=now, last_seen=now)
            row_num = ws_raw.max_row + 1
            ws_raw.append(row)

            # Row highlight based on no_website + priority
            no_website = not bool((biz.get("website") or "").strip())
            priority   = biz.get("priority", "low")
            fill = _row_fill(no_website, priority)

            if fill:
                for col_idx in range(1, len(COLUMNS) + 1):
                    ws_raw.cell(row=row_num, column=col_idx).fill = fill

            existing_ids.add(biz_id)
            existing_urls.add(maps_url)
            new_count += 1

        self._refresh_filtered_sheets(wb)
        self._ensure_dropdowns(wb)
        _safe_save(wb, MASTER_FILE)
        logger.info("Saved %d new | %d duplicates skipped → %s", new_count, dup_count, MASTER_FILE)
        return new_count, dup_count

    def refresh_views(self):
        """Rebuild no_website_leads and high_priority from raw_leads without adding data."""
        if not MASTER_FILE.exists():
            logger.error("Master file not found: %s", MASTER_FILE)
            return
        wb = self._load_or_create()
        self._ensure_sheets(wb)
        self._compact_raw(wb)
        self._refresh_filtered_sheets(wb)
        self._ensure_dropdowns(wb)
        _safe_save(wb, MASTER_FILE)
        logger.info("Filtered sheets refreshed.")

    def save_snapshot(self, businesses: list[dict]):
        """Write a standalone dated snapshot to data/runs/."""
        date_str = datetime.now().strftime("%Y_%m_%d_%H%M%S")
        path     = RUNS_DIR / f"scrape_{date_str}.xlsx"
        wb       = _create_new_workbook()
        ws       = wb["raw_leads"]
        now      = datetime.now().strftime("%Y-%m-%d")

        for biz in businesses:
            maps_url = normalize_url(biz.get("maps_url") or "")
            biz_id   = generate_business_id(
                maps_url=maps_url,
                name=biz.get("name", ""),
                address=biz.get("address", ""),
            )
            biz["business_id"] = biz_id
            biz["maps_url"]    = maps_url
            biz.update(calculate_score(biz))
            ws.append(_business_to_row(biz, first_seen=now, last_seen=now))

        _safe_save(wb, path)
        logger.info("Snapshot saved → %s", path)

    # ── Internal sheet refresh ─────────────────────────────────────────────────

    def _refresh_filtered_sheets(self, wb: Workbook):
        ws_raw = wb["raw_leads"]
        if ws_raw.max_row < 2:
            return

        headers = [ws_raw.cell(row=1, column=c).value for c in range(1, len(COLUMNS) + 1)]

        try:
            no_web_idx  = headers.index("no_website")
            priority_idx = headers.index("priority")
        except ValueError:
            logger.warning("Expected columns not found in raw_leads, skipping view refresh.")
            return

        no_web_rows   = []
        high_prio_rows = []

        for row in ws_raw.iter_rows(min_row=2, values_only=True):
            if row[no_web_idx]:
                no_web_rows.append(row)
                if row[priority_idx] == "high":
                    high_prio_rows.append(row)

        self._rebuild_sheet(wb, "no_website_leads", no_web_rows)
        self._rebuild_sheet(wb, "high_priority",    high_prio_rows)

    def _rebuild_sheet(self, wb: Workbook, name: str, rows: list):
        # Safety: never destroy protected sheets
        if name in _PROTECTED_SHEETS:
            return

        if name in wb.sheetnames:
            del wb[name]

        ws = wb.create_sheet(name)
        _write_headers(ws)

        for row in rows:
            ws.append(list(row))

        logger.debug("Rebuilt '%s' → %d rows", name, len(rows))
