"""
Corrige los links de Google Maps de los leads YA guardados.

Contexto: una versión vieja del sistema guardaba el maps_url normalizado
(sin el CID y en minúsculas), por lo que apuntaba a una búsqueda genérica y a
veces caía en otro negocio. El CID original ya no se puede recuperar, así que
reconstruimos un link confiable a partir de:
  1. nombre + coordenadas  (lo más preciso: centra el mapa en el negocio), o
  2. nombre + domicilio     (cuando no hay coordenadas).

A futuro ya no ocurre: el scraper ahora guarda el link real completo.

Uso:
    python -m src.fix_links              # corrige Excel y Supabase
    python -m src.fix_links --excel      # solo el Excel local
    python -m src.fix_links --supabase   # solo la base del dashboard
"""

import argparse
import logging
import re
from urllib.parse import quote_plus

import pandas as pd
from openpyxl import load_workbook

from config import MASTER_FILE
from src.excel_manager import _safe_save
from src.utils import setup_logging

logger = logging.getLogger(__name__)

_GLYPH_RE = re.compile("[-\U0001f000-\U0001faff]")


def _clean(s) -> str:
    """Quita glifos de íconos y saltos de línea de un texto."""
    if s is None:
        return ""
    return _GLYPH_RE.sub("", str(s)).replace("\n", " ").replace("\r", " ").strip()


def _valid_coord(x) -> bool:
    try:
        return x is not None and not pd.isna(x) and str(x).strip() != ""
    except Exception:
        return x is not None


def build_maps_url(name, address="", lat=None, lng=None) -> str:
    """Construye un link de Google Maps confiable hacia el negocio."""
    name = _clean(name)
    if _valid_coord(lat) and _valid_coord(lng):
        try:
            return (
                f"https://www.google.com/maps/search/"
                f"{quote_plus(name)}/@{float(lat)},{float(lng)},17z"
            )
        except (TypeError, ValueError):
            pass
    addr = _clean(address)
    query = f"{name}, {addr}" if addr else name
    return f"https://www.google.com/maps/search/?api=1&query={quote_plus(query)}"


def fix_excel() -> int:
    if not MASTER_FILE.exists():
        logger.error("No existe el Excel maestro: %s", MASTER_FILE)
        return 0
    wb = load_workbook(str(MASTER_FILE))
    ws = wb["raw_leads"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    try:
        col = {h: headers.index(h) + 1 for h in ("name", "address", "lat", "lng", "maps_url")}
    except ValueError as exc:
        logger.error("Faltan columnas en raw_leads: %s", exc)
        return 0

    n = 0
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=col["name"]).value
        if not name or not str(name).strip():
            continue
        addr = ws.cell(row=r, column=col["address"]).value
        lat = ws.cell(row=r, column=col["lat"]).value
        lng = ws.cell(row=r, column=col["lng"]).value
        ws.cell(row=r, column=col["maps_url"]).value = build_maps_url(name, addr, lat, lng)
        ws.cell(row=r, column=col["address"]).value = _clean(addr)  # limpia glifo de paso
        n += 1

    _safe_save(wb, MASTER_FILE)
    logger.info("Excel: %d links corregidos", n)
    return n


def fix_supabase() -> int:
    from src.db import SupabaseSync
    sync = SupabaseSync.from_env()
    if not sync:
        logger.warning("Supabase no configurado (sin .env) — omitido")
        return 0
    client = sync.client
    res = client.table("leads").select("business_id, name, address, lat, lng").execute()
    rows = res.data or []
    n = 0
    for row in rows:
        url = build_maps_url(row.get("name"), row.get("address"), row.get("lat"), row.get("lng"))
        client.table("leads").update(
            {"maps_url": url, "address": _clean(row.get("address"))}
        ).eq("business_id", row["business_id"]).execute()
        n += 1
    logger.info("Supabase: %d links corregidos", n)
    return n


def _cli():
    setup_logging()
    p = argparse.ArgumentParser(description="Corregir links de Google Maps existentes")
    p.add_argument("--excel", action="store_true", help="Solo el Excel local")
    p.add_argument("--supabase", action="store_true", help="Solo la base del dashboard")
    args = p.parse_args()

    do_both = not args.excel and not args.supabase
    if args.excel or do_both:
        fix_excel()
    if args.supabase or do_both:
        fix_supabase()


if __name__ == "__main__":
    _cli()
